import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Data extracted from the codebase
objects = [
    ("Codeunit", 73030575, "GSO Load Test Data - Inventory"),
    ("Codeunit", 73030576, "GSO ABC Analysis Job"),
    ("Codeunit", 73030577, "GSO ABC Install"),
    ("Codeunit", 73030578, "GSO ABC Upgrade"),
    ("Codeunit", 73030579, "GSO ABC Upgrade Helper"),
    ("Codeunit", 73030580, "GSO ABC Job Queue Helper"),
    ("Codeunit", 73030581, "GSO ABC Telemetry"),
    ("Codeunit", 73030582, "GSO ABC Analysis Mgt."),
    ("Codeunit", 73030583, "GSO ABC Worksheet Export"),
    ("Codeunit", 73030584, "GSO ABC Worksheet Stats"),
    ("Codeunit", 73030585, "GSO ABC Upgrade v2"),
    ("Codeunit", 73030586, "GSO ABC Upgrade Helper v2"),
    ("Enum", 73030575, "GSO ABC Class"),
    ("Enum", 73030576, "GSO ABC Model"),
    ("Page", 73030575, "GSO ABC Analysis Setup"),
    ("Page", 73030576, "GSO ABC Counting Policy List"),
    ("Page", 73030577, "GSO ABC Count Policy Subform"),
    ("Page", 73030578, "GSO ABC Counting Policy Card"),
    ("Page", 73030579, "GSO ABC Counting Policy Lines"),
    ("Page", 73030580, "GSO ABC Worksheet"),
    ("Page", 73030581, "GSO ABC Worksheet Params"),
    ("Page", 73030582, "GSO ABC Stats Part"),
    ("Page", 73030583, "GSO ABC History"),
    ("Page", 73030584, "GSO ABC History Lines"),
    ("Page", 73030585, "GSO ABC History Card"),
    ("Page", 73030586, "GSO Load Test Runner"),
    ("PageExtension", 73030575, "GSO Item Card Ext."),
    ("PageExtension", 73030576, "GSO Item List Ext."),
    ("PageExtension", 73030577, "GSO ABC Inventory Manager RC"),
    ("PageExtension", 73030578, "GSO SKU Card Ext."),
    ("PermissionSet", 73030575, "GSO ABC ANALYSIS"),
    ("PermissionSet", 73030576, "GSO ABC MANAGER"),
    ("PermissionSet", 73030577, "GSO ABC ADMIN"),
    ("PermissionSet", 73030578, "GSO ABC USER"),
    ("Report", 73030575, "GSO ABC Distribution"),
    ("Report", 73030576, "GSO ABC Item Movement"),
    ("Report", 73030577, "GSO ABC Analysis"),
    ("Report", 73030578, "GSO ABC Exceptions"),
    ("Report", 73030579, "GSO ABC Comparison"),
    ("Report", 73030580, "GSO ABC Detail by Class"),
    ("Table", 73030575, "GSO ABC Analysis Setup"),
    ("Table", 73030576, "GSO ABC Upgrade Tag"),
    ("Table", 73030577, "GSO ABC Counting Policy Header"),
    ("Table", 73030578, "GSO ABC Counting Policy Line"),
    ("Table", 73030579, "GSO ABC Preview Line"),
    ("Table", 73030580, "GSO ABC History Header"),
    ("Table", 73030581, "GSO ABC History Line"),
    ("Table", 73030582, "GSO ABC Stats"),
    ("TableExtension", 73030575, "GSO ABC Item Ext."),
    ("TableExtension", 73030576, "GSO ABC SKU Ext."),
]

# Sort objects by type and then by ID
objects.sort(key=lambda x: (x[0], x[1]))

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ABC Basic"

# Header styling
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=11)
header_alignment = Alignment(horizontal="center", vertical="center")

# Add headers
ws['A1'] = "Object Type"
ws['B1'] = "ID"
ws['C1'] = "Object Name"

for cell in ['A1', 'B1', 'C1']:
    ws[cell].fill = header_fill
    ws[cell].font = header_font
    ws[cell].alignment = header_alignment

# Add data
row = 2
for obj_type, obj_id, obj_name in objects:
    ws[f'A{row}'] = obj_type
    ws[f'B{row}'] = obj_id
    ws[f'C{row}'] = obj_name
    row += 1

# Adjust column widths
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 15
ws.column_dimensions['C'].width = 40

# Create summary worksheet
ws_summary = wb.create_sheet("Summary by Object Type")

# Add summary headers
ws_summary['A1'] = "Object Type"
ws_summary['B1'] = "Count"
ws_summary['C1'] = "First ID"
ws_summary['D1'] = "Last ID"

for cell in ['A1', 'B1', 'C1', 'D1']:
    ws_summary[cell].fill = header_fill
    ws_summary[cell].font = header_font
    ws_summary[cell].alignment = header_alignment

# Calculate summary by object type
summary = {}
for obj_type, obj_id, obj_name in objects:
    if obj_type not in summary:
        summary[obj_type] = {"count": 0, "ids": []}
    summary[obj_type]["count"] += 1
    summary[obj_type]["ids"].append(obj_id)

# Add summary data
row = 2
for obj_type in sorted(summary.keys()):
    ws_summary[f'A{row}'] = obj_type
    ws_summary[f'B{row}'] = summary[obj_type]["count"]
    ws_summary[f'C{row}'] = min(summary[obj_type]["ids"])
    ws_summary[f'D{row}'] = max(summary[obj_type]["ids"])
    row += 1

# Adjust column widths for summary
ws_summary.column_dimensions['A'].width = 20
ws_summary.column_dimensions['B'].width = 12
ws_summary.column_dimensions['C'].width = 15
ws_summary.column_dimensions['D'].width = 15

# Save the file
output_path = r"C:\GMSoft\OneDrive - Gmsoft\Applications\GMSOFT Apps.xlsx"
wb.save(output_path)
print(f"Excel file created successfully at: {output_path}")
print(f"Total objects: {len(objects)}")
print(f"Object types: {len(summary)}")
